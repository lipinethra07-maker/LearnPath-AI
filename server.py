import os
import time
import json
import io

from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS

from google import genai
from google.genai import types

from pypdf import PdfReader
from pptx import Presentation
from PIL import Image


# =====================================================
# FLASK SETUP
# =====================================================

frontend_folder = os.path.abspath(
    os.path.join(
        os.path.dirname(__file__),
        "../frontend"
    )
)

app = Flask(
    __name__,
    static_folder=frontend_folder,
    static_url_path=""
)

CORS(app)


# =====================================================
# GEMINI SETUP
# =====================================================

api_key = os.environ.get("GEMINI_API_KEY")

if not api_key:
    print("WARNING: GEMINI_API_KEY is not set.")
    client = None
else:
    client = genai.Client(api_key=api_key)

MODEL = "gemini-2.5-flash"


# =====================================================
# GEMINI HELPER
# =====================================================

def ask_gemini(prompt, image_bytes=None, require_json=False):

    if client is None:
        raise Exception("GEMINI_API_KEY is not configured.")

    max_attempts = 3

    for attempt in range(max_attempts):

        try:

            print(
                f"Sending request to Gemini... "
                f"(attempt {attempt + 1}/{max_attempts})"
            )

            contents = [prompt]

            if image_bytes:
                image = Image.open(
                    io.BytesIO(image_bytes)
                )
                contents.append(image)

            config = None

            if require_json:
                config = types.GenerateContentConfig(
                    response_mime_type="application/json"
                )

            response = client.models.generate_content(
                model=MODEL,
                contents=contents,
                config=config
            )

            if response and response.text:

                print("Gemini response received!")

                text = response.text.strip()

                if text.startswith("```"):
                    text = (
                        text
                        .replace("```json", "")
                        .replace("```", "")
                        .strip()
                    )

                return text

            raise Exception(
                "Gemini returned an empty response."
            )

        except Exception as e:

            error_message = str(e)

            print(
                "Gemini error:",
                error_message
            )

            temporary_error = (
                "503" in error_message
                or "UNAVAILABLE" in error_message
                or "429" in error_message
                or "500" in error_message
            )

            if temporary_error:

                if attempt < max_attempts - 1:

                    wait_time = 2 ** attempt

                    print(
                        f"Retrying in {wait_time} seconds..."
                    )

                    time.sleep(wait_time)

                else:

                    raise Exception(
                        "Gemini is temporarily unavailable."
                    )

            else:

                raise


# =====================================================
# FILE EXTRACTION
# =====================================================

def extract_file_text(file):

    filename = (
        file.filename or ""
    ).lower()

    print(
        "Reading file:",
        file.filename
    )

    try:

        if filename.endswith(".pdf"):

            reader = PdfReader(file)

            text = ""

            for page in reader.pages:

                page_text = page.extract_text()

                if page_text:
                    text += page_text + "\n"

            return text.strip()


        elif filename.endswith(".pptx"):

            presentation = Presentation(file)

            text = ""

            for slide in presentation.slides:

                for shape in slide.shapes:

                    if hasattr(shape, "text"):

                        if shape.text.strip():

                            text += (
                                shape.text +
                                "\n"
                            )

            return text.strip()


        elif filename.endswith(".txt"):

            return (
                file.read()
                .decode(
                    "utf-8",
                    errors="ignore"
                )
                .strip()
            )


        elif filename.endswith(
            (".png", ".jpg", ".jpeg", ".webp")
        ):

            return "[IMAGE_FILE]"


    except Exception as e:

        print(
            "Error extracting file:",
            str(e)
        )

        return None

    return ""


# =====================================================
# ANALYZE
# =====================================================

@app.route("/analyze", methods=["POST"])
def analyze():

    print("\n==============================")
    print("ANALYZE REQUEST RECEIVED")
    print("==============================")

    content = ""
    language = "English"
    image_bytes = None

    # ---------------------------------------------
    # FILE
    # ---------------------------------------------

    if "file" in request.files:

        file = request.files["file"]

        language = request.form.get(
            "language",
            "English"
        )

        if not file.filename:

            return jsonify({
                "error": "No file selected."
            }), 400

        extracted = extract_file_text(file)

        if extracted == "[IMAGE_FILE]":

            file.seek(0)

            image_bytes = file.read()

            content = (
                "Analyze the provided image study material."
            )

        elif extracted:

            content = extracted

        else:

            return jsonify({
                "error":
                    "Could not read the uploaded file."
            }), 400

    # ---------------------------------------------
    # TEXT / LINK
    # ---------------------------------------------

    else:

        data = request.get_json(
            silent=True
        )

        if data:

            content = str(
                data.get(
                    "content",
                    ""
                )
            ).strip()

            language = data.get(
                "language",
                "English"
            )

    if not content and not image_bytes:

        return jsonify({
            "error":
                "No readable study material found."
        }), 400

    print(
        "Extracted text characters:",
        len(content)
    )

    prompt = f"""
You are LearnPath AI, a personal learning assistant.

Analyze the study material below.

STUDY MATERIAL:
{content}

Use {language}.

Return ONLY valid JSON:

{{
  "topic": "Identify the main topic",
  "explanation": "Explain the material in very easy student-friendly language",
  "summary": "Give a short summary of the provided material",
  "hint": "Give quick revision hints useful before an examination",
  "points": [
    "Point 1",
    "Point 2",
    "Point 3",
    "Point 4",
    "Point 5"
  ],
  "questions": [
    "Practice Question 1",
    "Practice Question 2",
    "Practice Question 3"
  ]
}}

Rules:

- Analyze ANY subject or topic.
- Do not assume the material is about Machine Learning.
- Use ONLY the provided material/image.
- Do not invent unrelated information.
- Keep the explanation easy for a student.
- Return exactly 5 important points.
- Return exactly 3 practice questions.
- Use {language}.
- Return ONLY JSON.
"""

    try:

        result = ask_gemini(
            prompt,
            image_bytes=image_bytes,
            require_json=True
        )

        analysis = json.loads(result)

        return jsonify({
            "topic":
                analysis.get("topic", ""),

            "explanation":
                analysis.get("explanation", ""),

            "summary":
                analysis.get("summary", ""),

            "hint":
                analysis.get("hint", ""),

            "points":
                analysis.get("points", []),

            "questions":
                analysis.get("questions", []),

            "content":
                content
        })

    except json.JSONDecodeError:

        print(
            "Analysis returned invalid JSON."
        )

        return jsonify({
            "error":
                "AI returned an invalid format. Please try again."
        }), 500

    except Exception as e:

        print(
            "ANALYZE ERROR:",
            str(e)
        )

        return jsonify({
            "error":
                "AI is temporarily unavailable."
        }), 503


# =====================================================
# GENERATE QUIZ
# =====================================================

@app.route("/generate-quiz", methods=["POST"])
def generate_quiz():

    data = request.get_json(
        silent=True
    )

    if not data:

        return jsonify({
            "error":
                "No content received."
        }), 400

    content = str(
        data.get(
            "content",
            ""
        )
    ).strip()

    language = data.get(
        "language",
        "English"
    )

    if not content:

        return jsonify({
            "error":
                "No study material provided."
        }), 400

    prompt = f"""
You are LearnPath AI.

Create a 5-question multiple-choice quiz ONLY from the study material below.

STUDY MATERIAL:
{content}

Use {language}.

Return ONLY valid JSON:

[
  {{
    "question": "Question here",
    "options": [
      "Option A",
      "Option B",
      "Option C",
      "Option D"
    ],
    "answer": "Option A"
  }}
]

Rules:

- Create exactly 5 questions.
- Each question must have exactly 4 options.
- The answer must exactly match one option.
- Use ONLY information from the material.
- Questions can be from ANY subject.
- Do not assume the subject is Machine Learning.
- Use simple student-friendly language.
- Use {language}.
- Return ONLY JSON.
"""

    try:

        result = ask_gemini(
            prompt,
            require_json=True
        )

        quiz = json.loads(result)

        if not isinstance(quiz, list):

            raise ValueError(
                "Invalid quiz format."
            )

        valid_quiz = []

        for item in quiz[:5]:

            if not isinstance(
                item,
                dict
            ):
                continue

            question = item.get(
                "question"
            )

            options = item.get(
                "options"
            )

            answer = item.get(
                "answer"
            )

            if (
                question
                and isinstance(options, list)
                and len(options) == 4
                and answer in options
            ):

                valid_quiz.append({
                    "question":
                        question,

                    "options":
                        options,

                    "answer":
                        answer
                })

        if not valid_quiz:

            return jsonify({
                "error":
                    "AI returned an invalid quiz format."
            }), 500

        return jsonify({
            "quiz":
                valid_quiz
        })

    except json.JSONDecodeError:

        return jsonify({
            "error":
                "AI returned an invalid quiz format."
        }), 500

    except Exception as e:

        print(
            "QUIZ ERROR:",
            str(e)
        )

        return jsonify({
            "error":
                "AI is temporarily unavailable."
        }), 503


# =====================================================
# CHAT
# =====================================================

@app.route("/chat", methods=["POST"])
def chat():

    data = request.get_json(
        silent=True
    )

    if not data:

        return jsonify({
            "error":
                "No content received."
        }), 400

    message = str(
        data.get(
            "message",
            ""
        )
    ).strip()

    context = str(
        data.get(
            "context",
            ""
        )
    ).strip()

    is_revision = bool(
        data.get(
            "isRevision",
            False
        )
    )

    language = data.get(
        "language",
        "English"
    )

    if (
        not message
        and not is_revision
    ):

        return jsonify({
            "error":
                "No message provided."
        }), 400

    if is_revision:

        prompt = f"""
You are LearnPath AI, a friendly study companion.

The student wants a Last Minute Revision.

Based ONLY on the study material below, provide:

- Important concepts
- Important definitions
- Key points
- Quick revision hints
- 2 quick MCQs

Keep it short and easy to scan.

Use {language}.

STUDY MATERIAL:
{context}

Do not include unrelated information.
"""

    else:

        prompt = f"""
You are LearnPath AI, a friendly study companion.

Answer the student's question based ONLY on the provided study material.

If the answer is not present in the material,
politely say that it is not available in the uploaded material.

Use {language}.

STUDY MATERIAL:
{context}

STUDENT QUESTION:
{message}

Keep the answer clear and student-friendly.
"""

    try:

        result = ask_gemini(
            prompt
        )

        resources = []

        if is_revision:

            resources = [
                {
                    "title":
                        "Quick Revision",

                    "desc":
                        "Review important concepts from your study material."
                }
            ]

        return jsonify({

            "response":
                result,

            "resources":
                resources

        })

    except Exception as e:

        print(
            "CHAT ERROR:",
            str(e)
        )

        return jsonify({
            "error":
                "AI chat is temporarily unavailable."
        }), 503


# =====================================================
# LEARNING PATH
# =====================================================

@app.route(
    "/learning-path",
    methods=["POST"]
)
def learning_path():

    data = request.get_json(
        silent=True
    )

    if not data:

        return jsonify({
            "error":
                "No content received."
        }), 400

    content = str(
        data.get(
            "content",
            ""
        )
    ).strip()

    language = data.get(
        "language",
        "English"
    )

    if not content:

        return jsonify({
            "error":
                "No study material provided."
        }), 400

    prompt = f"""
You are LearnPath AI.

Create a personalized learning path from this study material.

STUDY MATERIAL:
{content}

Use {language}.

Return ONLY valid JSON:

{{
  "topic": "Main topic",

  "learning_path": [
    {{
      "step": 1,
      "title": "Step title",
      "description": "Simple description"
    }},
    {{
      "step": 2,
      "title": "Step title",
      "description": "Simple description"
    }},
    {{
      "step": 3,
      "title": "Step title",
      "description": "Simple description"
    }},
    {{
      "step": 4,
      "title": "Step title",
      "description": "Simple description"
    }},
    {{
      "step": 5,
      "title": "Step title",
      "description": "Simple description"
    }}
  ]
}}

Rules:

- Use ONLY the provided study material.
- Do not assume a particular subject.
- Create exactly 5 steps.
- Keep descriptions simple.
- Use {language}.
- Return ONLY JSON.
"""

    try:

        result = ask_gemini(
            prompt,
            require_json=True
        )

        learning_data =
            json.loads(result)

        return jsonify(
            learning_data
        )

    except json.JSONDecodeError:

        return jsonify({
            "error":
                "AI returned an invalid learning path format."
        }), 500

    except Exception as e:

        print(
            "LEARNING PATH ERROR:",
            str(e)
        )

        return jsonify({
            "error":
                "AI is temporarily unavailable."
        }), 503


# =====================================================
# HOME
# =====================================================

@app.route("/")
def home():

    return send_from_directory(
        app.static_folder,
        "index.html"
    )


# =====================================================
# HEALTH CHECK
# =====================================================

@app.route(
    "/health",
    methods=["GET"]
)
def health():

    return jsonify({

        "status":
            "running",

        "service":
            "LearnPath AI",

        "model":
            MODEL,

        "gemini_configured":
            client is not None

    })


# =====================================================
# START SERVER
# =====================================================

if __name__ == "__main__":

    print("")
    print("==============================")
    print("       LEARNPATH AI SERVER")
    print("==============================")
    print("Model:", MODEL)
    print(
        "Gemini configured:",
        client is not None
    )
    print(
        "Frontend:",
        frontend_folder
    )
    print(
        "Server:",
        "http://127.0.0.1:5000"
    )
    print("==============================")
    print("")

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=False
    )