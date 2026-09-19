import os
import time
import json

from flask import Flask, request, jsonify
from flask_cors import CORS
from google import genai
from pypdf import PdfReader
from pptx import Presentation

app = Flask(__name__)
CORS(app)

# ===============================
# GEMINI SETUP
# ===============================

api_key = os.environ.get("GEMINI_API_KEY")

if not api_key:
    raise RuntimeError("GEMINI_API_KEY is not set.")

client = genai.Client(api_key=api_key)

# Use the model that worked earlier
MODEL = "gemini-3.5-flash"


# ===============================
# GEMINI HELPER
# ===============================

def ask_gemini(prompt):

    max_attempts = 3

    for attempt in range(max_attempts):

        try:

            print(
                f"Sending request to Gemini... "
                f"(attempt {attempt + 1}/{max_attempts})"
            )

            response = client.models.generate_content(
                model=MODEL,
                contents=prompt
            )

            if response and response.text:

                print("Gemini response received!")

                return response.text.strip()

            raise Exception("Gemini returned an empty response.")

        except Exception as e:

            error_message = str(e)

            print("Gemini error:", error_message)

            if (
                "503" in error_message
                or "UNAVAILABLE" in error_message
                or "429" in error_message
                or "500" in error_message
            ):

                if attempt < max_attempts - 1:

                    wait_time = 2 ** attempt

                    print(
                        f"Temporary Gemini error. "
                        f"Retrying in {wait_time} seconds..."
                    )

                    time.sleep(wait_time)

                else:

                    raise Exception(
                        "Gemini is temporarily unavailable."
                    )

            else:

                raise


# ===============================
# EXTRACT FILE TEXT
# ===============================

def extract_file_text(file):

    filename = file.filename.lower()

    print("Reading file:", file.filename)

    # PDF
    if filename.endswith(".pdf"):

        reader = PdfReader(file)

        text = ""

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:
                text += page_text

            text += "\n"

        return text

    # PPTX
    elif filename.endswith(".pptx"):

        presentation = Presentation(file)

        text = ""

        for slide in presentation.slides:

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text += shape.text
                    text += "\n"

        return text

    # TXT
    elif filename.endswith(".txt"):

        return file.read().decode(
            "utf-8",
            errors="ignore"
        )

    return ""


# ===============================
# ANALYZE
# ===============================

@app.route("/analyze", methods=["POST"])
def analyze():

    print("\n==============================")
    print("ANALYZE REQUEST RECEIVED")
    print("==============================")

    content = ""

    # Uploaded file
    if "file" in request.files:

        print("Input type: FILE")

        file = request.files["file"]

        print("File name:", file.filename)

        content = extract_file_text(file)

    # Text input
    else:

        print("Input type:", request.content_type)

        data = request.get_json(silent=True)

        if data:

            content = data.get("content", "")

    # Check content
    if not content or not content.strip():

        print("No readable content found.")

        return jsonify({
            "result": "No readable study material found."
        }), 400

    print("Extracted characters:", len(content))

    # Prompt
    prompt = f"""
You are LearnPath AI, a personal learning assistant.

Analyze the study material below.

STUDY MATERIAL:
{content}

Use simple English.

Give the following:

1. Topic
2. Simple Explanation
3. Short Summary
4. Last-Minute Hint
5. 5 Important Points
6. 3 Practice Questions
7. Suggested Next Topic

Important rules:

- Use ONLY the information from the provided study material.
- Do not add unrelated information.
- Do not assume the subject.
- Identify the actual topic from the material.
- Keep the explanation easy for a student to understand.
"""

    try:

        result = ask_gemini(prompt)

        return jsonify({
            "result": result
        })

    except Exception as e:

        print("ANALYZE ERROR:", str(e))

        return jsonify({
            "result":
            "AI is temporarily unavailable. "
            "Please try Analyze again."
        }), 503


# ===============================
# GENERATE QUIZ
# ===============================

@app.route("/generate-quiz", methods=["POST"])
def generate_quiz():

    data = request.get_json(silent=True)

    if not data:

        return jsonify({
            "error": "No content received"
        }), 400

    content = data.get(
        "content",
        ""
    ).strip()

    if not content:

        return jsonify({
            "error": "No study material provided"
        }), 400

    prompt = f"""
You are LearnPath AI.

Create a 5-question multiple-choice quiz ONLY from
the study material below.

STUDY MATERIAL:

{content}

Return ONLY valid JSON.

Format:

[
  {{
    "question": "Question here",
    "options": [
      "Option A",
      "Option B",
      "Option C",
      "Option D"
    ],
    "answer": "Option A"
  }}
]

Rules:

- Create exactly 5 questions.
- Each question must have exactly 4 options.
- Use ONLY information from the provided material.
- Do not introduce unrelated topics.
- Keep questions suitable for students.
"""

    try:

        result = ask_gemini(prompt)

        if result.startswith("```"):

            result = result.replace("```json", "")
            result = result.replace("```", "")
            result = result.strip()

        quiz = json.loads(result)

        return jsonify({
            "quiz": quiz
        })

    except json.JSONDecodeError:

        print("Quiz returned invalid JSON.")

        return jsonify({
            "error":
            "AI returned an invalid quiz format. "
            "Please try again."
        }), 500

    except Exception as e:

        print("QUIZ ERROR:", str(e))

        return jsonify({
            "error":
            "AI is temporarily unavailable. "
            "Please try again."
        }), 503


# ===============================
# HOME
# ===============================

@app.route("/")
def home():

    return "LearnPath AI server is running!"


# ===============================
# START SERVER
# ===============================

if __name__ == "__main__":

    print("")
    print("==============================")
    print("LEARNPATH AI SERVER")
    print("==============================")
    print("Model:", MODEL)
    print("Server starting...")
    print("")

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=False
    )