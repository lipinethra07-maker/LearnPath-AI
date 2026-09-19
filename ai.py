import os
from flask import Flask, request, jsonify
from flask_cors import CORS
from google import genai
from pypdf import PdfReader
from pptx import Presentation

app = Flask(__name__)
CORS(app)

# ===============================
# GEMINI SETUP
# ===============================

client = genai.Client(
    api_key=os.environ["GEMINI_API_KEY"]
)


# ===============================
# FILE TEXT EXTRACTION
# ===============================

def extract_file_text(file):

    filename = file.filename.lower()

    # PDF
    if filename.endswith(".pdf"):

        reader = PdfReader(file)

        text = ""

        for page in reader.pages:
            text += page.extract_text() or ""

        return text


    # PPTX
    elif filename.endswith(".pptx"):

        presentation = Presentation(file)

        text = ""

        for slide in presentation.slides:

            for shape in slide.shapes:

                if hasattr(shape, "text"):
                    text += shape.text + "\n"

        return text


    # TXT
    elif filename.endswith(".txt"):

        return file.read().decode("utf-8", errors="ignore")


    return ""


# ===============================
# AI ANALYZE
# ===============================

@app.route("/analyze", methods=["POST"])
def analyze():

    content = ""

    # ===========================
    # TEXT FROM WEBSITE
    # ===========================

    if request.is_json:

        data = request.get_json(silent=True) or {}

        content = data.get("content", "")


    # ===========================
    # FILE / FORM DATA
    # ===========================

    else:

        if request.form.get("content"):

            content = request.form.get("content")


        elif "file" in request.files:

            file = request.files["file"]

            content = extract_file_text(file)


    # ===========================
    # CHECK CONTENT
    # ===========================

    if not content.strip():

        return jsonify({
            "result": "No study material found in this file."
        })


    # ===========================
    # GEMINI PROMPT
    # ===========================

    prompt = f"""
You are LearnPath AI, a personal learning assistant.

Analyze the following study material:

-------------------------
STUDY MATERIAL
-------------------------

{content}

-------------------------

Give the answer in simple English.

Include:

1. Topic
2. Simple Explanation
3. Short Summary
4. Last-Minute Hint
5. 5 Important Points
6. 3 Practice Questions

Keep the explanation useful for a student.
Do not add information that is unrelated to the study material.
"""


    # ===========================
    # GEMINI RESPONSE
    # ===========================

    response = client.models.generate_content(
        model="gemini-3.5-flash",
        contents=prompt
    )


    # ===========================
    # SEND RESULT TO WEBSITE
    # ===========================

    return jsonify({
        "result": response.text
    })


# ===============================
# HOME
# ===============================

@app.route("/")
def home():

    return "LearnPath AI server is running!"


# ===============================
# START SERVER
# ===============================

if __name__ == "__main__":

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=True
    )